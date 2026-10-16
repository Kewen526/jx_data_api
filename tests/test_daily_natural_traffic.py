"""日报自然流量计算及实际 Excel 导出回归测试。"""
import importlib.util
import sys
import tempfile
import types
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl


class DailyNaturalTrafficTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        config = types.ModuleType('app.core.config')
        config.settings = types.SimpleNamespace(TEMP_DIR=self.temp_dir.name)
        database = types.ModuleType('app.core.database')
        database.get_db_pool = lambda: None
        spec = importlib.util.spec_from_file_location(
            'report_under_test', Path(__file__).parents[1] / 'app/services/report.py'
        )
        self.report = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {'app.core.config': config, 'app.core.database': database}):
            spec.loader.exec_module(self.report)
        self.addCleanup(self.temp_dir.cleanup)

    def test_calculation_boundaries(self):
        cases = [
            ((314, Decimal('969.47'), Decimal('4.55')), 100.93),
            ((20, Decimal('0'), None), 20),
            ((20, 0, 0), 20),
            ((0, 0, None), 0),
            ((276, Decimal('1074.76'), Decimal('2.31')), -189.26),
            ((10, 1, 8), 9.88),
            ((20, 1, None), '数据异常'),
            ((20, 1, 0), '数据异常'),
            ((20, 1, -1), '数据异常'),
            ((None, 0, None), '数据缺失'),
            ((20, None, 1), '数据缺失'),
            ((20, -1, 1), '数据异常'),
            ((20, 1, 'NaN'), '数据异常'),
            ((20, 1, 'invalid'), '数据异常'),
        ]
        for args, expected in cases:
            with self.subTest(args=args):
                self.assertEqual(self.report.calculate_natural_traffic(*args), expected)

    def test_export_positions_values_and_formatting(self):
        base = dict.fromkeys([
            'shop_name', 'exposure_users', 'order_users', 'verify_users',
            'order_coupon_count', 'verify_coupon_count', 'phone_clicks',
            'address_clicks', 'new_good_review_count', 'intent_rate',
            'order_sale_amount', 'verify_sale_amount', 'verify_after_discount',
            'order_user_rank', 'verify_amount_rank', 'new_review_count',
            'new_collect_users', 'is_force_offline', 'consult_users',
            'checkin_count', 'promotion_order_count', 'ad_balance',
        ], 0)
        base.update(shop_name='测试门店', order_user_rank=2, verify_amount_rank=3)
        inputs = [(314, Decimal('969.47'), Decimal('4.55')),
                  (20, 0, None), (276, Decimal('1074.76'), Decimal('2.31')),
                  (20, 1, 0), (20, None, None)]
        rows = [dict(base, shop_id=i, visit_users=v, promotion_cost=c,
                     click_avg_price=p) for i, (v, c, p) in enumerate(inputs, 1)]

        class Cursor:
            def execute(self, *args): pass
            def fetchall(self): return rows
            def close(self): pass

        class Connection:
            def cursor(self, **kwargs): return Cursor()
            def close(self): pass

        db = types.SimpleNamespace(get_connection=Connection)
        with patch.object(self.report, 'get_db_pool', return_value=db), \
             patch.object(self.report, 'get_shop_info_mapping', return_value={}), \
             patch.object(self.report, 'get_region_info_mapping', return_value={}), \
             patch.object(self.report, 'get_coupon_orders_last_7days', return_value=0), \
             patch.object(self.report, 'get_ad_orders_today', return_value=0):
            output = self.report.generate_daily_report('2026-10-15')

        wb = openpyxl.load_workbook(output, data_only=True)
        self.addCleanup(wb.close)
        summary = wb['汇总']
        self.assertEqual(summary.max_column, 24)
        self.assertEqual([summary[f'{c}1'].value for c in 'UVWX'],
                         ['优惠后核销金额', '自然流量', '下单人数商圈排名', '核销金额商圈排名'])
        expected = [100.93, 20, -189.26, '数据异常', '数据缺失']
        for i, (sheet, value) in enumerate(zip(list(wb)[1:], expected), 2):
            with self.subTest(sheet=sheet.title):
                self.assertEqual(sheet.max_row, 42)
                self.assertEqual(sheet['A36'].value, '自然流量：')
                self.assertEqual(sheet['B36'].value, value)
                self.assertEqual(summary[f'V{i}'].value, value)
                self.assertEqual(sheet['B36'].number_format, '0.00')
                self.assertEqual(summary[f'V{i}'].number_format, '0.00')
                self.assertEqual(sheet['B5'].value, inputs[i-2][0])
                self.assertEqual(sheet['A20'].value, '推广通消耗：')
                self.assertEqual(sheet['A21'].value, '推广通点击单价：')
                self.assertEqual(sheet['A38'].value, '团单被强制下线数量：')
                self.assertEqual(sheet['A40'].value, '运营：')
                self.assertEqual(summary[f'W{i}'].value, '第2名')
                self.assertEqual(summary[f'X{i}'].value, '第3名')
                self.assertEqual(summary[f'X{i}'].border.right.style, 'thin')
                self.assertEqual(sheet['C42'].border.bottom.style, 'thin')


if __name__ == '__main__':
    unittest.main()
